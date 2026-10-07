"""The workout log as a spreadsheet (xlsx), one row per logged exercise."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import WeightUnit, WorkoutExerciseLog, WorkoutLog
from photo_helpers import other_client

DAY = date(2026, 9, 14)


@pytest.fixture
def logged(me):
    with SessionLocal() as s:
        log = WorkoutLog(owner_id=me, day_label="Push day", plan_name="Export Plan", log_date=DAY)
        s.add(log)
        s.flush()
        s.add(WorkoutExerciseLog(workout_log_id=log.id, name="Bench Press", completed=True, sets=3, reps_value=10, weight_value=135.0, weight_unit=WeightUnit.LB,
                                 gross_kcal=60.0, net_kcal=45.0, volume_lb=4050.0, met=3.5))
        s.add(WorkoutExerciseLog(workout_log_id=log.id, name="Push-up", completed=False, sets=2, reps_value=15))
        s.commit()
        lid = log.id
    yield lid
    with SessionLocal() as s:
        s.query(WorkoutLog).filter(WorkoutLog.plan_name == "Export Plan").delete()
        s.commit()


def sheet(client):
    r = client.get("/workouts/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"] and 'attachment; filename="amide-workout-log.xlsx"' == r.headers["content-disposition"]
    return list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))


def test_one_row_per_logged_exercise_with_a_header(client, db, logged):
    rows = sheet(client)
    assert rows[0][:6] == ("Date", "Plan", "Workout", "Exercise", "Done", "Sets")
    mine = [r for r in rows[1:] if r[1] == "Export Plan"]
    assert len(mine) == 2
    bench = next(r for r in mine if r[3] == "Bench Press")
    assert bench[0] == "2026-09-14" and bench[2] == "Push day" and bench[4] == "Yes" and bench[5] == 3 and bench[6] == 10 and bench[7] == 135.0 and bench[8] == "lb"
    assert 45.0 in bench and 4050.0 in bench
    assert next(r for r in mine if r[3] == "Push-up")[4] == "No"


def test_only_my_own_workouts_are_exported(client, db, logged):
    with other_client("exportother") as member:
        rows = list(load_workbook(io.BytesIO(member.get("/workouts/export.xlsx").content)).active.iter_rows(values_only=True))
    assert all(r[1] != "Export Plan" for r in rows[1:])


def test_the_workouts_page_links_to_the_export(client, db):
    assert 'href="/workouts/export.xlsx"' in client.get("/workouts").text
