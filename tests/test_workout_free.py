"""A workout that is not part of any plan: you name it and add the exercises on the day."""

import html
from datetime import date

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import BodyMeasurement, WorkoutExerciseLog, WorkoutLog
from photo_helpers import other_client

DAY = date.today().isoformat()


@pytest.fixture(autouse=True)
def _clean(client, me):
    def wipe():
        with SessionLocal() as s:
            s.query(WorkoutLog).filter(WorkoutLog.owner_id == me, WorkoutLog.plan_day_id.is_(None)).delete()
            s.query(BodyMeasurement).filter(BodyMeasurement.owner_id == me).delete()
            s.commit()
    wipe()
    with SessionLocal() as s:
        s.add(BodyMeasurement(owner_id=me, measured_at=date.today(), weight_lbs=180.0))
        s.commit()
    yield
    wipe()


def post(client, **extra):
    data = {"log_date": DAY, "label": "Park workout", "extra-0-exercise": "Back Squat", "extra-0-sets_value": "3", "extra-0-reps_value": "10"} | extra
    return client.post("/workouts/free/log", data=data, follow_redirects=False)


def free_logs(me):
    with SessionLocal() as s:
        return [(log.id, log.day_label, log.plan_day_id, [e.name for e in log.exercise_logs]) for log in s.scalars(
            select(WorkoutLog).where(WorkoutLog.owner_id == me, WorkoutLog.plan_day_id.is_(None)))]


def test_the_free_workout_form_asks_for_a_name_and_the_exercises(client, db):
    page = html.unescape(client.get("/workouts/free/log").text)
    assert 'action="/workouts/free/log"' in page and 'name="label"' in page and "Add exercise" in page


def test_a_free_workout_is_saved_with_no_plan_day_and_a_calorie_estimate(client, db, me):
    r = post(client)
    assert r.status_code == 303
    [(log_id, label, day_id, names)] = free_logs(me)
    assert label == "Park workout" and day_id is None and names == ["Back Squat"]
    with SessionLocal() as s:
        assert s.scalar(select(WorkoutExerciseLog.net_kcal).where(WorkoutExerciseLog.workout_log_id == log_id)) > 0


def test_it_shows_in_the_journal_and_the_export(client, db, me):
    post(client)
    assert "Park workout" in client.get("/measurements", params={"tab": "journal"}).text
    import io
    from openpyxl import load_workbook
    rows = list(load_workbook(io.BytesIO(client.get("/workouts/export.xlsx").content)).active.iter_rows(values_only=True))
    assert any(r[2] == "Park workout" and r[3] == "Back Squat" for r in rows[1:])


def test_a_workout_needs_a_name_and_at_least_one_exercise(client, db, me):
    assert post(client, label="  ").status_code == 422
    assert client.post("/workouts/free/log", data={"log_date": DAY, "label": "Empty"}).status_code == 422
    assert free_logs(me) == []


def test_an_exercise_that_is_not_in_the_database_is_refused_like_any_other(client, db, me):
    assert post(client, **{"extra-0-exercise": "Zzz Quasar Lift"}).status_code == 422


def test_an_earlier_free_workout_can_be_reopened_and_replaced(client, db, me):
    post(client)
    [(log_id, *_)] = free_logs(me)
    page = html.unescape(client.get(f"/workouts/free/{log_id}/log").text)
    assert "Park workout" in page and "Back Squat" in page and f'action="/workouts/free/{log_id}/log"' in page
    client.post(f"/workouts/free/{log_id}/log", data={"log_date": DAY, "label": "Renamed", "extra-0-exercise": "Push-up", "extra-0-sets_value": "2", "extra-0-reps_value": "15"})
    assert [(label, [n.lower() for n in names]) for _, label, _, names in free_logs(me)] == [("Renamed", ["push-up"])]


def test_someone_elses_free_workout_cannot_be_opened_or_replaced(client, db, me):
    post(client)
    [(log_id, *_)] = free_logs(me)
    with other_client("freeother") as member:
        assert member.get(f"/workouts/free/{log_id}/log").status_code == 404
        assert member.post(f"/workouts/free/{log_id}/log", data={"log_date": DAY, "label": "x", "extra-0-exercise": "Back Squat"}).status_code == 404


def test_the_workouts_page_and_the_journal_link_to_it(client, db):
    assert 'href="/workouts/free/log"' in client.get("/workouts").text
    assert "/workouts/free/log" in client.get("/measurements", params={"tab": "journal"}).text
