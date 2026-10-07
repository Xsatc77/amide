"""Reading received files into the importer's own `PriceListData`, whatever they are: photos, spreadsheets, typed text, PDFs."""

import io
import re
from pathlib import Path

from app import config
from app.library.price_lists import ocr
from app.library.price_lists.names import propagate_names
from app.library.price_lists.reader import (
    ParsedRow, PriceListData, read_pdf, rows_from_lines, rows_from_table, scan_notes, unread_spec_lines,
)

_MAX_SHEETS, _MAX_ROWS, _MAX_COLS = 10, 5000, 40
MAX_UNCOMPRESSED = 100 * 1024 * 1024
_PRICE_WORDS = re.compile(r"price|quote|catalog", re.IGNORECASE)


class ReadError(ValueError):
    """The file cannot be read at all; the message is short and safe to show."""


class OcrMissing(ReadError):
    """Pictures of text were given but the text-recognition engine is not installed."""


def _data(rows: list[ParsedRow], lines: list[str]) -> PriceListData:
    note, hint = scan_notes(lines)
    return PriceListData(rows=rows, shipping_note=note, warehouse_hint=hint, unread_spec_lines=unread_spec_lines(lines, rows))


def has_price(row: ParsedRow) -> bool:
    """A real price: a stock sheet's quantities (0 in stock) are not prices."""
    return row.pack_price is not None and row.pack_price > 0


def prefer_priced(table_rows: list[ParsedRow], line_rows: list[ParsedRow]) -> list[ParsedRow]:
    """Of a table reading and a line-by-line reading of the same page, the one that found more prices (the table wins a tie)."""
    def priced(rows):
        return sum(1 for r in rows if has_price(r))
    return line_rows if priced(line_rows) > priced(table_rows) else table_rows


def read_pdf_file(path: Path, recognize=None) -> PriceListData:
    try:
        return read_pdf(path, recognize=recognize)
    except ocr.OcrUnavailable:
        raise OcrMissing("scanned pages need text recognition, which is not installed") from None
    except Exception as exc:
        raise ReadError(f"the PDF could not be read ({type(exc).__name__})") from None


def read_images(images, recognize=None) -> PriceListData:
    """Photos of a price list, read as the pages of one list."""
    images = list(images)
    if len(images) > config.INGEST_MAX_IMAGES:
        raise ReadError("too many photos for one list")
    rows: list[ParsedRow] = []
    lines: list[str] = []
    for number, image in enumerate(images, 1):
        if image.width * image.height > config.PHOTO_MAX_PIXELS:
            raise ReadError("a photo is too large")
        try:
            words = (recognize or ocr.recognize)(image)
        except ocr.OcrUnavailable:
            raise OcrMissing("photos need text recognition, which is not installed") from None
        page_lines, page_rows = ocr.rows_from_words(words, number)
        lines.extend(page_lines)
        rows.extend(prefer_priced(page_rows, rows_from_lines(page_lines, number)))
    return _data(rows, lines)


def read_xlsx(data: bytes) -> PriceListData:
    try:
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise ReadError("the spreadsheet is too large")
    except ReadError:
        raise
    except Exception:
        raise ReadError("the spreadsheet could not be opened") from None
    try:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ReadError("the spreadsheet could not be opened") from None
    rows: list[ParsedRow] = []
    lines: list[str] = []
    try:
        for number, sheet in enumerate(book.worksheets[:_MAX_SHEETS], 1):
            table = []
            for raw in sheet.iter_rows(min_row=1, max_row=_MAX_ROWS, max_col=_MAX_COLS, values_only=True):
                cells = ["" if v is None else (f"{v:g}" if isinstance(v, float) else str(v)).strip() for v in raw]
                if any(cells):
                    table.append(cells)
                    lines.append(" ".join(c for c in cells if c))
            sheet_rows = rows_from_table(table, number) if table else []
            rows.extend(sheet_rows or rows_from_lines([" ".join(c for c in r if c) for r in table], number))
    except Exception:
        raise ReadError("the spreadsheet could not be read") from None
    finally:
        book.close()
    propagate_names(rows)
    return _data(rows, lines)


def read_text(text: str) -> PriceListData:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return _data(rows_from_lines(lines), lines)


def priced_fraction(data: PriceListData) -> float:
    return (sum(1 for r in data.rows if has_price(r)) / len(data.rows)) if data.rows else 0.0


def looks_like_price_list(data: PriceListData, caption: str | None, filename: str | None) -> tuple[bool, str | None]:
    """A price list has at least 3 rows with a price (and most rows priced), or says "price" in its name or caption and
    has at least one priced row and most rows priced. Anything else is chatter, a photo of something else, a certificate,
    or a warehouse stock sheet (its quantities read as prices, mostly zero)."""
    priced = sum(1 for r in data.rows if has_price(r))
    named = bool(_PRICE_WORDS.search(caption or "") or _PRICE_WORDS.search(filename or ""))
    if priced >= 3 and priced_fraction(data) >= 0.6:
        return True, None
    if named and priced >= 1 and priced_fraction(data) >= 0.5:
        return True, None
    return False, "not a price list (too few priced rows)"
